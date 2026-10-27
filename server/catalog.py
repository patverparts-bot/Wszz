"""Parse private price files and project the catalogue for a trusted user record."""
import io
import math
import re
import zipfile
from datetime import date, datetime
from openpyxl import load_workbook


def norm(value):
    return re.sub(r'[^a-z0-9\u0531-\u058f]', '', str(value or '').lower())


def number(value):
    try:
        result = float(value or 0)
        if not math.isfinite(result):
            raise ValueError()
        return result
    except (TypeError, ValueError):
        raise ValueError('Некорректное числовое значение в прайсе')


def read_rows(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 3000 or sum(i.file_size for i in archive.infolist()) > 80_000_000:
                raise ValueError('Слишком большой распакованный файл')
            if any('vbaProject' in i.filename for i in archive.infolist()):
                raise ValueError('Макросы запрещены')
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        sheet = book.worksheets[0]
        if sheet.max_row > 50000 or sheet.max_column > 150:
            raise ValueError('Допустимо до 50 000 строк и 150 колонок')
        rows = []
        for cells in sheet:
            if any(c.data_type == 'f' for c in cells):
                raise ValueError('Сохраните значения вместо формул перед загрузкой')
            rows.append([c.value.isoformat() if isinstance(c.value, (datetime, date)) else c.value for c in cells])
        book.close()
        return rows
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Не удалось прочитать XLSX') from exc


def parse_price(raw):
    rows = read_rows(raw)
    header_index = next((i for i, row in enumerate(rows[:25]) if any(
        any(term in str(c or '').lower() for term in ['artikul', 'артикул', 'արտիկուլ']) for c in row)), None)
    if header_index is None:
        raise ValueError('Не найдена колонка артикула')
    headers = [str(c or '').strip() for c in rows[header_index]]
    def find(*terms):
        return next((i for i, h in enumerate(headers) if any(t.lower() in h.lower() for t in terms)), None)
    cols = dict(art=find('art', 'артикул', 'արտիկուլ'), brand=find('brand', 'бренд', 'բրենդ'),
                category=find('group', 'kategoriya', 'группа', 'категория', 'խումբ', 'կատեգորիա'),
                name=find('name', 'naimenovanie', 'наименование', 'номенклатура', 'товар', 'անվանում'),
                retail=find('Զեղչված', 'discount'), base=find('roznich', 'մանրածախ', 'розница'),
                ws=find('wholesale', 'опт', 'մեծածախ'), kpPromo=find("Key Partner's promo AMD"),
                kp=find("Key Partner's AMD"))
    if any(cols[k] is None for k in ['art', 'brand', 'category', 'name', 'ws']):
        raise ValueError('Нужны колонки: артикул, бренд, категория, наименование, Wholesale')
    if cols['retail'] is None:
        cols['retail'] = cols['base']
    if cols['retail'] is None:
        raise ValueError('Не найдена розничная цена')
    stock_cols = [(i, h) for i, h in enumerate(headers) if re.fullmatch(r'\d{3}', h)]
    if not stock_cols:
        raise ValueError('Не найдены колонки складов (003, 023 и т. п.)')
    products, seen = [], set()
    for row_no, row in enumerate(rows[header_index + 1:], header_index + 2):
        def get(key):
            idx = cols[key]
            return row[idx] if idx is not None and idx < len(row) else None
        art = str(get('art') or '').strip()
        if not art:
            continue
        if art in seen:
            raise ValueError(f'Повтор артикула в строке {row_no}')
        if len(art) > 30:
            raise ValueError(f'Слишком длинный артикул в строке {row_no}')
        seen.add(art)
        prices = {key: number(get(key)) for key in ['retail', 'base', 'ws', 'kp', 'kpPromo']}
        if any(v < 0 for v in prices.values()):
            raise ValueError(f'Отрицательная цена в строке {row_no}')
        stocks = {name: number(row[i] if i < len(row) else 0) for i, name in stock_cols}
        if any(v < 0 for v in stocks.values()):
            raise ValueError(f'Отрицательный остаток в строке {row_no}')
        strings = {k: str(get(k) or '').strip() for k in ['brand', 'category', 'name']}
        if any(len(v) > 1000 for v in strings.values()):
            raise ValueError(f'Слишком длинное текстовое поле в строке {row_no}')
        products.append(dict(articul=art, **strings, **prices, stocks=stocks))
    if not products:
        raise ValueError('Прайс не содержит товаров')
    return products


def project(products, user):
    result = []
    admin = user['role'] == 'admin'
    kp = user['sheet'] == 'KP' and not admin
    domus = any(term in user['client'].upper() for term in ['ԴՈՄՈՒՍ', 'ДОМУՍ', 'DOMUS'])
    for p in products:
        brand = norm(p['brand'])
        if kp and (not user['allBrands'] and brand not in [norm(b) for b in user['brands']]):
            continue
        if kp and (brand == 'miele' or (brand == 'sony' and norm('Հեռուստացույց') in norm(p['category']))):
            continue
        wholesale = (p['kpPromo'] or p['kp'] or p['ws']) if kp else p['ws']
        kp_price = p['kpPromo'] or p['kp']
        retail = p['retail']
        if not admin and (wholesale <= 0 or (retail > 0 and wholesale > retail)):
            continue
        base_stock = sum(q for n, q in p['stocks'].items() if n in ['003', '023', '026', '031'])
        other_stock = sum(q for n, q in p['stocks'].items() if n not in ['003', '023', '026', '031'])
        if admin:
            if base_stock + other_stock <= 0:
                continue
            stock = ', '.join(f'{n}: {q:g}' for n, q in p['stocks'].items() if q > 0)
        else:
            if base_stock == 0 and other_stock < 5:
                continue
            stock = 'Առկա է' if base_stock >= 3 else 'Ճշտել'
        def margin(price):
            return (retail / price if domus else (1 - price / retail) * 100) if price > 0 and retail > 0 else 0
        item = {k: p[k] for k in ['articul', 'brand', 'category', 'name']}
        item.update(retail=retail, wholesale=wholesale, isPromo=retail > 0 and p['base'] > retail,
                    isWholesalePromo=bool(kp and p['kpPromo']), hasWholesale=wholesale > 0 or (admin and kp_price > 0),
                    marginVal=margin(wholesale), stockStatus=stock)
        if admin:
            item.update(kpPrice=kp_price, kpIsPromo=bool(p['kpPromo']), kpMarginVal=margin(kp_price), targetWarehousesStock=base_stock)
        if admin or user.get('allowResale'):
            item['resaleWholesale'] = p['ws']
        result.append(item)
    return result
