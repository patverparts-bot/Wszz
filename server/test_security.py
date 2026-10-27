import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from werkzeug.security import generate_password_hash
from server.app import create_app


def price_file(ws=80):
    book = Workbook()
    sheet = book.active
    sheet.append(['Артикул','Brand','Group','Name','Զեղչված գին','Մանրածախ գին','Wholesale AMD',"Key Partner's AMD","Key Partner's promo AMD",'003','025'])
    sheet.append(['1','LG','TV','LG TV',100,120,ws,70,65,4,0])
    sheet.append(['2','Miele','Home','Miele',100,100,80,70,0,4,0])
    sheet.append(['3','Sony','Հեռուստացույց','Sony TV',100,100,80,70,0,4,0])
    sheet.append(['4','Samsung','TV','Samsung',100,100,80,70,0,4,0])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class SecurityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.admin_hash = generate_password_hash('Admin-password-123',method='scrypt')
        cls.client_hash = generate_password_hash('Client-password-123',method='scrypt')

    def setUp(self):
        scratch = Path(__file__).resolve().parents[1] / 'work' / 'tests'
        scratch.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=scratch)
        self.app = create_app(self.temp.name, testing=True)
        self.client = self.app.test_client()
        self.database = Path(self.temp.name)/'zzws.sqlite'
        with sqlite3.connect(self.database) as db:
            for uid, role, sheet in [('admin','admin','Dealers'),('dealer','client','Dealers'),('kp','client','KP')]:
                profile = dict(role=role,sheet=sheet,client=uid,allBrands=False,brands=['LG'],rowsList=[],allowResale=False)
                db.execute('INSERT INTO users VALUES(?,?,?,?,1)',(uid,uid,self.admin_hash if uid=='admin' else self.client_hash,json.dumps(profile)))
        db.close()
        self.headers = {'X-ZZWS-Request':'1','Origin':'http://127.0.0.1:8080'}

    def tearDown(self):
        self.temp.cleanup()

    def post(self, path, **kwargs):
        return self.client.post('/api/'+path,headers=self.headers,**kwargs)

    def login(self, who='admin'):
        return self.post('login',json=dict(login=who,password='Admin-password-123' if who=='admin' else 'Client-password-123'))

    def publish(self, raw=None):
        response = self.post('admin/price/preview',data={'file':(io.BytesIO(raw or price_file()),'price.xlsx')})
        self.assertEqual(response.status_code,200,response.json)
        return self.post('admin/price/publish',json={'draft':response.json['draft']})

    def test_anonymous_and_static_files(self):
        for path in ['/api/me','/api/catalog','/api/catalog/export','/api/admin/users','/api/admin/audit']:
            self.assertEqual(self.client.get(path).status_code,401,path)
        for path in ['/pass.xlsx','/PriceList_3.xlsx','/private-data/zzws.sqlite','/server/app.py','/.git/config','/backups/archive.zip','/js/../../pass.xlsx']:
            self.assertEqual(self.client.get(path).status_code,404,path)
        for path in ['/', '/admin']:
            response = self.client.get(path)
            self.assertEqual(response.status_code,200)
            response.close()

    def test_csrf_and_cookie(self):
        self.assertEqual(self.client.post('/api/login',json={}).status_code,403)
        self.assertEqual(self.client.post('/api/login',headers={**self.headers,'Origin':'https://evil.example'},json={}).status_code,403)
        response=self.login()
        self.assertEqual(response.status_code,200)
        cookie=response.headers['Set-Cookie']
        self.assertIn('HttpOnly',cookie)
        self.assertIn('SameSite=Strict',cookie)
        self.assertEqual(self.client.get('/api/me').json['role'],'admin')
        self.post('logout')
        self.assertEqual(self.client.get('/api/catalog').status_code,401)

    def test_prices_brands_and_role_tampering(self):
        self.login();self.publish();self.post('logout');self.login('dealer')
        data=self.client.get('/api/catalog?role=admin&sheet=KP').json
        self.assertEqual(data['products'][0]['wholesale'],80)
        for p in data['products']:
            for key in ['kpPrice','kp','kpPromo','resaleWholesale','stocks','targetWarehousesStock']:
                self.assertNotIn(key,p)
        self.assertEqual(self.client.get('/api/admin/users').status_code,403)
        self.assertEqual(self.post('admin/price/publish',json={'draft':'fake'}).status_code,403)
        self.post('logout');self.login('kp')
        products=self.client.get('/api/catalog').json['products']
        self.assertEqual([p['brand'] for p in products],['LG'])
        self.assertEqual(products[0]['wholesale'],65)
        self.assertNotIn('ws',products[0])

    def test_preview_publish_rollback_and_invalid_file(self):
        self.login()
        first=self.publish().json['version']
        response=self.post('admin/price/preview',data={'file':(io.BytesIO(price_file(90)),'p.xlsx')})
        self.assertEqual(self.client.get('/api/catalog').json['products'][0]['wholesale'],80)
        self.post('admin/price/publish',json={'draft':response.json['draft']})
        self.assertEqual(self.client.get('/api/catalog').json['products'][0]['wholesale'],90)
        invalid=self.post('admin/price/preview',data={'file':(io.BytesIO(b'bad'),'p.xlsx')})
        self.assertEqual(invalid.status_code,400)
        self.assertEqual(self.client.get('/api/catalog').json['products'][0]['wholesale'],90)
        self.post('admin/price/rollback',json={'version':first})
        self.assertEqual(self.client.get('/api/catalog').json['products'][0]['wholesale'],80)

    def test_block_and_reset_revoke_sessions(self):
        self.login(); user_client=self.app.test_client()
        user_client.post('/api/login',headers=self.headers,json={'login':'dealer','password':'Client-password-123'})
        payload=dict(id='dealer',login='dealer',client='Dealer',sheet='Dealers',role='client',active=False,brands=[])
        self.assertEqual(self.post('admin/users',json=payload).status_code,200)
        self.assertEqual(user_client.get('/api/me').status_code,401)
        payload.update(active=True,password='New-password-123')
        self.post('admin/users',json=payload)
        self.assertEqual(user_client.post('/api/login',headers=self.headers,json={'login':'dealer','password':'Client-password-123'}).status_code,401)
        self.assertEqual(user_client.post('/api/login',headers=self.headers,json={'login':'dealer','password':'New-password-123'}).status_code,200)

    def test_export_contains_only_own_prices(self):
        self.login();self.publish();self.post('logout');self.login('kp')
        response=self.client.get('/api/catalog/export')
        self.assertEqual(response.status_code,200)
        rows=list(load_workbook(io.BytesIO(response.data),read_only=True).active.values)
        self.assertEqual(len(rows),3)
        self.assertEqual(rows[2][5],65)
        self.assertNotIn('KP',rows[1])

    def test_lockout(self):
        for _ in range(10):
            self.assertEqual(self.post('login',json=dict(login='missing',password='bad')).status_code,401)
        self.assertEqual(self.login().status_code,429)

    def test_formula_rejected(self):
        book=load_workbook(io.BytesIO(price_file()))
        book.active['E2']='=1+2'
        output=io.BytesIO();book.save(output)
        self.login()
        self.assertEqual(self.post('admin/price/preview',data={'file':(io.BytesIO(output.getvalue()),'p.xlsx')}).status_code,400)

    def test_user_import_preview_no_deletions_or_password_reuse(self):
        self.login()
        book=Workbook()
        book.active.append(['login','client','sheet','brands','allBrands','organization','address','contact'])
        book.active.append(['dealer','Updated dealer','KP','LG','no','','',''])
        book.active.append(['new-client','New client','Dealers','','no','','',''])
        output=io.BytesIO();book.save(output)
        preview=self.post('admin/users/preview',data={'file':(io.BytesIO(output.getvalue()),'users.xlsx')})
        self.assertEqual(preview.status_code,200,preview.json)
        self.assertEqual(len(self.client.get('/api/admin/users').json['users']),3)
        self.assertEqual(self.post('admin/users/import',json={'draft':preview.json['draft']}).status_code,200)
        users={u['login']:u for u in self.client.get('/api/admin/users').json['users']}
        self.assertEqual(len(users),4)
        self.assertFalse(users['new-client']['active'])
        self.assertEqual(users['dealer']['sheet'],'KP')
        self.assertEqual(users['admin']['role'],'admin')
        self.post('logout')
        self.assertEqual(self.login('dealer').status_code,200)


if __name__ == '__main__':
    unittest.main()
